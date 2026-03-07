# main.py
import re
import html
import logging

from fastapi import FastAPI, HTTPException, UploadFile, File, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, field_validator
from google import genai
import tempfile
import subprocess
import shutil
import glob
import os
import json
import docx

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="LaTeX Studio API", version="2.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Initialize Gemini — keep your own key here
client = genai.Client(api_key="API_KEY_HERE")

GEMINI_MODEL = "gemini-2.5-flash"

# ─── Pydantic Models ─────────────────────────────────────────────────────────

class EditorBlock(BaseModel):
    type: str
    data: dict

class Author(BaseModel):
    name: str
    affiliation: str = ""

class CompileRequest(BaseModel):
    template: str
    title: str
    author: str = ""          # legacy single-author (kept for compat)
    affiliation: str = ""     # legacy affiliation
    authors: list[Author] = [] # new multi-author list
    blocks: list[EditorBlock]

    @field_validator("template")
    @classmethod
    def validate_template(cls, v: str) -> str:
        allowed = {"standard", "acl", "cvpr"}
        if v not in allowed:
            raise ValueError(f"template must be one of {allowed}")
        return v

    @field_validator("title", "author", "affiliation")
    @classmethod
    def strip_fields(cls, v: str) -> str:
        return v.strip()

    @field_validator("blocks")
    @classmethod
    def validate_blocks(cls, v: list) -> list:
        if len(v) > 500:
            raise ValueError("Too many blocks (max 500)")
        return v

class RephraseRequest(BaseModel):
    text: str
    style: str = "academic"

class ReviewRequest(BaseModel):
    template: str
    title: str
    authors: list[Author] = []
    blocks: list[EditorBlock]


# ─── LaTeX Sanitisation ───────────────────────────────────────────────────────

_LATEX_ESCAPES: list[tuple[str, str]] = [
    ("\\", "\\textbackslash{}"),
    ("&",  "\\&"),
    ("%",  "\\%"),
    ("$",  "\\$"),
    ("#",  "\\#"),
    ("_",  "\\_"),
    ("{",  "\\{"),
    ("}",  "\\}"),
    ("~",  "\\textasciitilde{}"),
    ("^",  "\\textasciicircum{}"),
]

_UNICODE_REPLACEMENTS: dict[str, str] = {
    "\u201c": "``", "\u201d": "''", "\u2018": "`", "\u2019": "'",
    "\u2013": "--", "\u2014": "---", "\u2012": "-", "\u2015": "---",
    "\u00a0": " ", "\u202f": " ", "\u2009": " ", "\u2008": " ",
    "\u2007": " ", "\u2006": " ", "\u2005": " ", "\u2004": " ",
    "\u2003": " ", "\u2002": " ", "\u2001": " ", "\u2000": " ",
    "\u200b": "", "\u200c": "", "\u200d": "", "\u200e": "",
    "\u200f": "", "\ufeff": "", "\u00ad": "",
    "\u2026": "\\dots", "\u2212": "-", "\u00d7": "\\times",
    "\u00f7": "\\div", "\u2032": "'", "\u2033": "''",
    "\u2192": "\\rightarrow{}", "\u2190": "\\leftarrow{}",
    "\u2194": "\\leftrightarrow{}", "\u21d2": "\\Rightarrow{}",
    "\u2264": "\\leq{}", "\u2265": "\\geq{}", "\u2260": "\\neq{}",
    "\u221e": "\\infty{}", "\u03b1": "\\alpha{}", "\u03b2": "\\beta{}",
    "\u0394": "\\Delta{}", "\u223c": "\\sim{}", "\u03b3": "\\gamma{}",
    "\u03b4": "\\delta{}", "\u03bb": "\\lambda{}", "\u03bc": "\\mu{}",
    "\u03c3": "\\sigma{}", "\u03c4": "\\tau{}", "\u03c0": "\\pi{}",
    "\u03a9": "\\Omega{}",
}


def sanitize_latex(text: str) -> str:
    for char, replacement in _UNICODE_REPLACEMENTS.items():
        text = text.replace(char, replacement)
    for char, replacement in _LATEX_ESCAPES:
        text = text.replace(char, replacement)
    return text


def clean_html(raw: str) -> str:
    raw = raw.replace("&nbsp;", " ")
    raw = html.unescape(raw)
    raw = re.sub(r"<br\s*/?>", "\n", raw)
    raw = re.sub(r"<[^>]+>", "", raw)
    return raw.strip()


# ─── Author Formatting ────────────────────────────────────────────────────────

def _resolve_authors(request: CompileRequest) -> list[Author]:
    """Return list[Author] whether caller sent new multi-author or legacy fields."""
    if request.authors:
        return request.authors
    # Fall back to legacy single-author fields
    if request.author:
        return [Author(name=request.author, affiliation=request.affiliation)]
    return [Author(name="Author", affiliation="")]


def format_authors_latex(authors: list[Author], template: str) -> str:
    if not authors:
        return "Author"

    if template == "acl":
        parts = []
        for a in authors:
            name = sanitize_latex(a.name)
            affil = sanitize_latex(a.affiliation)
            parts.append(f"{name} \\\\\n\t{affil}" if affil else name)
        return "\n\t\\And\n\t".join(parts)

    elif template == "cvpr":
        names = " \\quad ".join(sanitize_latex(a.name) for a in authors)
        # Deduplicate affiliations while preserving order
        seen, affils = set(), []
        for a in authors:
            if a.affiliation and a.affiliation not in seen:
                seen.add(a.affiliation)
                affils.append(sanitize_latex(a.affiliation))
        affil_str = " \\quad ".join(affils)
        return f"{names} \\\\\n{affil_str}" if affil_str else names

    else:  # standard
        parts = []
        for a in authors:
            name = sanitize_latex(a.name)
            affil = sanitize_latex(a.affiliation)
            parts.append(f"{name} \\\\ \\small {affil}" if affil else name)
        return "\n\\and\n".join(parts)


# ─── Templates ───────────────────────────────────────────────────────────────

ACL_TEMPLATE = r"""\documentclass[11pt]{article}
\usepackage{acl}
\usepackage{times}
\usepackage{latexsym}
\usepackage{cite}
\usepackage{amsmath,amssymb}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage{booktabs}
\usepackage{array}
\usepackage{longtable}
\usepackage{microtype}
\begin{document}

\title{[[TITLE]]}
\author{
	[[AUTHORS]]
}
\maketitle
\begin{abstract}
	Generated abstract from editor content.
\end{abstract}

%%BODY%%

\bibliographystyle{acl_natbib}
\end{document}
"""

CVPR_TEMPLATE = r"""\documentclass[10pt,twocolumn,letterpaper]{article}
\usepackage{cvpr}
\usepackage{times}
\usepackage{epsfig}
\usepackage{graphicx}
\usepackage{amsmath,amssymb}
\usepackage{cite}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage{booktabs}
\usepackage{array}
\usepackage{longtable}
\begin{document}

\title{[[TITLE]]}
\author{[[AUTHORS]]}
\maketitle

%%BODY%%

\end{document}
"""

STANDARD_TEMPLATE = r"""\documentclass[12pt]{article}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage{amsmath,amssymb}
\usepackage{booktabs}
\usepackage{array}
\usepackage{longtable}
\usepackage{microtype}
\usepackage{geometry}
\geometry{margin=1in}
\title{[[TITLE]]}
\author{[[AUTHORS]]}
\date{\today}
\begin{document}
\maketitle
%%BODY%%
\end{document}
"""

TEMPLATES = {
    "acl": ACL_TEMPLATE,
    "cvpr": CVPR_TEMPLATE,
    "standard": STANDARD_TEMPLATE,
}

# ─── Template requirement definitions ────────────────────────────────────────

TEMPLATE_REQUIREMENTS = {
    "acl": {
        "name": "ACL",
        "required": ["abstract", "introduction", "limitations"],
        "recommended": ["related work", "methodology", "experiments", "results", "conclusion", "ethical considerations"],
    },
    "cvpr": {
        "name": "CVPR",
        "required": ["abstract", "introduction", "compute budget"],
        "recommended": ["related work", "method", "experiments", "results", "conclusion", "appendix"],
    },
    "standard": {
        "name": "Standard Article",
        "required": ["introduction"],
        "recommended": ["related work", "methodology", "results", "conclusion"],
    },
}


# ─── DOCX helpers ─────────────────────────────────────────────────────────────

def extract_text_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    return "\n\n".join(
        para.text for para in doc.paragraphs if para.text.strip()
    )


# ─── Chunked-conversion helpers ───────────────────────────────────────────────

CHUNK_CHARS = 7_000
CONTEXT_BLOCKS = 2

_BLOCK_RULES = """\
Block conversion rules (STRICT):
- "header"    → section / subsection headings  (level: 1 = section, 2 = subsection, 3 = subsubsection)
- "paragraph" → regular prose
- "code"      → any LaTeX math, equation, table, or verbatim environment already in LaTeX
- "list"      → bullet / numbered lists  (style: "unordered" or "ordered"; items: array of strings)

RETURN ONLY VALID JSON — no markdown fences, no commentary outside the JSON object."""

_FIRST_SCHEMA = """\
{
  "title": "string",
  "author": "string",
  "affiliation": "string",
  "blocks": [ {"type": "paragraph", "data": {"text": "..."}} ]
}"""

_CONT_SCHEMA = """\
{
  "blocks": [ {"type": "paragraph", "data": {"text": "..."}} ]
}"""


def _split_chunks(text: str, size: int = CHUNK_CHARS) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r'\n{2,}', text) if p.strip()]
    chunks: list[str] = []
    buf: list[str] = []
    buf_len = 0
    for para in paragraphs:
        para_len = len(para) + 2
        if buf and buf_len + para_len > size:
            chunks.append("\n\n".join(buf))
            buf, buf_len = [para], para_len
        else:
            buf.append(para)
            buf_len += para_len
    if buf:
        chunks.append("\n\n".join(buf))
    return chunks


def _call_gemini(prompt: str) -> dict:
    try:
        resp = client.models.generate_content(model=GEMINI_MODEL, contents=prompt)
        raw = resp.text.strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```\s*$", "", raw)
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        logger.error("Gemini non-JSON (%s): %.300s", exc, raw)
        raise HTTPException(status_code=500,
            detail=f"AI returned invalid JSON on chunk: {exc}")
    except Exception as exc:
        logger.error("Gemini error: %s", exc)
        raise HTTPException(status_code=500, detail=str(exc))


def _process_chunks(chunks: list[str]) -> dict:
    total = len(chunks)
    logger.info("Processing document in %d chunk(s)", total)

    first_prompt = f"""\
You are converting an academic document into Editor.js JSON blocks.
This is chunk 1 of {total}.

{_BLOCK_RULES}

SCHEMA:
{_FIRST_SCHEMA}

Document text (chunk 1/{total}):
{chunks[0]}
"""
    result = _call_gemini(first_prompt)
    all_blocks: list[dict] = result.get("blocks", [])
    metadata = {
        "title":       result.get("title", ""),
        "author":      result.get("author", ""),
        "affiliation": result.get("affiliation", ""),
    }
    logger.info("Chunk 1/%d → %d blocks, metadata=%s", total, len(all_blocks), metadata)

    for idx, chunk_text in enumerate(chunks[1:], start=2):
        context_tail = all_blocks[-CONTEXT_BLOCKS:] if len(all_blocks) >= CONTEXT_BLOCKS else all_blocks
        context_json = json.dumps(context_tail, ensure_ascii=False)

        cont_prompt = f"""\
You are converting an academic document into Editor.js JSON blocks.
This is chunk {idx} of {total} — continue from where the previous chunk ended.

The last {len(context_tail)} block(s) already produced (for continuity — do NOT repeat them):
{context_json}

{_BLOCK_RULES}

SCHEMA for this response (blocks only — metadata already captured):
{_CONT_SCHEMA}

Document text (chunk {idx}/{total}):
{chunk_text}
"""
        cont_result = _call_gemini(cont_prompt)
        new_blocks = cont_result.get("blocks", [])
        all_blocks.extend(new_blocks)
        logger.info("Chunk %d/%d → +%d blocks (total %d)", idx, total, len(new_blocks), len(all_blocks))

    return {**metadata, "blocks": all_blocks, "chunks_processed": total}


# ─── LaTeX construction ───────────────────────────────────────────────────────

def _list_items_to_latex(items: list, ordered: bool, depth: int = 0) -> str:
    env = "enumerate" if ordered else "itemize"
    indent = "  " * depth
    lines = [f"{indent}\\begin{{{env}}}"]
    for item in items:
        if isinstance(item, dict):
            text = sanitize_latex(clean_html(item.get("content", "")))
            lines.append(f"{indent}  \\item {text}")
            nested = item.get("items", [])
            if nested:
                lines.append(_list_items_to_latex(nested, ordered, depth + 1))
        else:
            lines.append(f"{indent}  \\item {sanitize_latex(clean_html(str(item)))}")
    lines.append(f"{indent}\\end{{{env}}}")
    return "\n".join(lines)


def _table_to_latex(table_data: dict) -> str:
    content = table_data.get("content", [])
    with_headings: bool = table_data.get("withHeadings", False)
    rows = [r for r in content if isinstance(r, list)]
    if not rows:
        return ""
    cols = max(len(r) for r in rows)
    if cols == 0:
        return ""
    col_spec = " ".join(["l"] * cols)
    lines = [
        "", "\\begin{table}[htbp]", "  \\centering",
        f"  \\begin{{tabular}}{{{col_spec}}}", "  \\toprule",
    ]
    for i, row in enumerate(rows):
        padded = list(row) + [""] * (cols - len(row))
        if i == 0 and with_headings:
            cells = " & ".join(
                "\\textbf{" + sanitize_latex(clean_html(str(c))) + "}" for c in padded
            )
        else:
            cells = " & ".join(sanitize_latex(clean_html(str(c))) for c in padded)
        lines.append(f"  {cells} \\\\")
        if i == 0 and with_headings:
            lines.append("  \\midrule")
    lines += ["  \\bottomrule", "  \\end{tabular}", "\\end{table}", ""]
    return "\n".join(lines)


def construct_latex(request: CompileRequest) -> str:
    body_parts: list[str] = []

    for block in request.blocks:
        btype = block.type
        data = block.data

        if btype == "paragraph":
            raw = data.get("text", "")
            text = sanitize_latex(clean_html(raw))
            if text:
                body_parts.append(f"{text}\n")

        elif btype == "header":
            level = int(data.get("level", 2))
            text = sanitize_latex(clean_html(data.get("text", "")))
            cmd = {1: "section", 2: "subsection", 3: "subsubsection"}.get(level, "subsubsection")
            if text:
                body_parts.append(f"\\{cmd}{{{text}}}\n")

        elif btype == "code":
            code = data.get("code", "").strip()
            if code:
                body_parts.append(f"{code}\n")

        elif btype == "list":
            style = data.get("style", "unordered")
            ordered = style == "ordered"
            items = data.get("items", [])
            if items:
                body_parts.append(_list_items_to_latex(items, ordered) + "\n")

        elif btype == "table":
            tex_table = _table_to_latex(data)
            if tex_table:
                body_parts.append(tex_table + "\n")

        else:
            logger.warning("Unsupported block type ignored: %s", btype)

    body = "\n".join(body_parts)
    authors = _resolve_authors(request)
    authors_latex = format_authors_latex(authors, request.template)

    tpl = TEMPLATES.get(request.template, STANDARD_TEMPLATE)
    result = (
        tpl
        .replace("[[TITLE]]",   sanitize_latex(request.title))
        .replace("[[AUTHORS]]", authors_latex)
        # Legacy placeholders (keep in case someone still uses old templates)
        .replace("[[AUTHOR]]",      sanitize_latex(authors[0].name if authors else ""))
        .replace("[[AFFILIATION]]", sanitize_latex(authors[0].affiliation if authors else ""))
        .replace("%%BODY%%",    body)
    )
    return result


# ─── Routes ───────────────────────────────────────────────────────────────────

@app.post("/api/upload")
async def upload_and_modularize(file: UploadFile = File(...)):
    if not (file.filename or "").endswith(".docx"):
        raise HTTPException(status_code=400, detail="Only .docx files are supported.")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name

    try:
        raw_text = extract_text_from_docx(tmp_path)
    except Exception as e:
        logger.error("Failed to extract docx: %s", e)
        raise HTTPException(status_code=422, detail=f"Could not parse document: {e}")
    finally:
        os.remove(tmp_path)

    raw_text = raw_text.strip()
    if not raw_text:
        raise HTTPException(status_code=422, detail="The uploaded document appears to be empty.")

    chunks = _split_chunks(raw_text)
    logger.info("Document %.0f chars → %d chunk(s)", len(raw_text), len(chunks))

    import asyncio
    loop = asyncio.get_event_loop()
    result = await loop.run_in_executor(None, _process_chunks, chunks)
    return result


@app.post("/api/latex")
def get_latex_source(request: CompileRequest) -> dict:
    try:
        return {"source": construct_latex(request)}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/compile")
def compile_latex(request: CompileRequest):
    tex_content = construct_latex(request)
    logger.info("Compiling %.0f chars of LaTeX (template=%s)", len(tex_content), request.template)

    with tempfile.TemporaryDirectory() as temp_dir:
        tex_path = os.path.join(temp_dir, "main.tex")
        pdf_path = os.path.join(temp_dir, "main.pdf")

        current_dir = os.getcwd()
        for ext in ("*.sty", "*.cls", "*.bib", "*.bst"):
            for fp in glob.glob(os.path.join(current_dir, ext)):
                shutil.copy(fp, temp_dir)

        with open(tex_path, "w", encoding="utf-8") as f:
            f.write(tex_content)

        latex_cmd = ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "main.tex"]

        def run_pass(pass_num: int):
            try:
                return subprocess.run(
                    latex_cmd, cwd=temp_dir, check=True,
                    stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=60,
                )
            except subprocess.TimeoutExpired:
                raise HTTPException(status_code=504,
                    detail=f"LaTeX compilation timed out (pass {pass_num}).")
            except subprocess.CalledProcessError as e:
                log_output = e.stdout.decode("utf-8", errors="ignore")
                log_lines = log_output.splitlines()
                error_chunks: list[str] = []
                for idx, line in enumerate(log_lines):
                    if line.startswith("!"):
                        chunk = log_lines[idx: idx + 5]
                        error_chunks.append("\n".join(chunk))
                if error_chunks:
                    short_error = "\n\n".join(error_chunks)
                else:
                    short_error = "\n".join(log_lines[-20:]) or "Unknown LaTeX error."
                raise HTTPException(status_code=400,
                    detail=f"LaTeX Compiler Error (pass {pass_num}):\n{short_error}")

        run_pass(1)
        run_pass(2)

        if not os.path.exists(pdf_path):
            raise HTTPException(status_code=500, detail="PDF generation failed silently.")

        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()

    logger.info("PDF compiled successfully (%.0f bytes)", len(pdf_bytes))
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"X-PDF-Size": str(len(pdf_bytes))},
    )


@app.post("/api/rephrase")
def rephrase_text(req: RephraseRequest):
    """Rephrase a selected passage using Gemini."""
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="No text provided.")

    style_map = {
        "academic":  "Rephrase the following text in a formal, academic style suitable for a research paper. Improve clarity and precision.",
        "simpler":   "Simplify the following text. Use plain language and shorter sentences while preserving the core meaning.",
        "concise":   "Make the following text more concise. Remove redundancy and tighten the writing without losing meaning.",
        "elaborate": "Expand and elaborate on the following text. Add nuance, detail, and academic depth while staying on topic.",
        "fluent":    "Improve the fluency and flow of the following text. Fix awkward phrasing while preserving the original meaning.",
    }
    instruction = style_map.get(req.style, style_map["academic"])

    prompt = f"""{instruction}

Return ONLY the rephrased text — no explanations, no quotation marks, no preamble.

Text:
{req.text}
"""
    try:
        resp = client.models.generate_content(model=GEMINI_MODEL, contents=prompt)
        rephrased = resp.text.strip().strip('"').strip("'")
        return {"rephrased": rephrased, "original": req.text}
    except Exception as e:
        logger.error("Rephrase error: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/review")
def review_paper(req: ReviewRequest):
    """Analyze a paper draft and return scores + section compliance."""
    reqs = TEMPLATE_REQUIREMENTS.get(req.template, TEMPLATE_REQUIREMENTS["standard"])

    # Build readable text and collect section headers
    text_parts: list[str] = []
    section_headers: list[str] = []
    word_count = 0

    for block in req.blocks:
        if block.type == "header":
            t = clean_html(block.data.get("text", ""))
            text_parts.append(f"\n## {t}\n")
            section_headers.append(t.lower().strip())
        elif block.type == "paragraph":
            t = clean_html(block.data.get("text", ""))
            text_parts.append(t)
            word_count += len(t.split())
        elif block.type == "list":
            for item in block.data.get("items", []):
                t = clean_html(item.get("content", "") if isinstance(item, dict) else str(item))
                text_parts.append("• " + t)
                word_count += len(t.split())

    paper_text = "\n".join(text_parts)

    all_sections = reqs["required"] + reqs["recommended"]

    prompt = f"""You are a senior academic reviewer for a {reqs['name']} paper submission.

Paper title: {req.title}
Template: {reqs['name']}
Word count (approx): {word_count}
Section headers found: {json.dumps(section_headers)}

REQUIRED sections for {reqs['name']}: {json.dumps(reqs['required'])}
RECOMMENDED sections: {json.dumps(reqs['recommended'])}

Paper content (first 5000 chars):
{paper_text[:5000]}

Evaluate the paper and return ONLY a valid JSON object (no markdown fences):
{{
  "section_check": {{
    "abstract":        {{"present": true,  "quality": "good",    "note": "Well-structured abstract covering motivation and results."}},
    "introduction":    {{"present": false, "quality": "missing", "note": "No introduction section found."}},
    "limitations":     {{"present": false, "quality": "missing", "note": "Required for ACL — must discuss model/data limitations."}}
  }},
  "scores": {{
    "completeness":    {{"score": 6, "comment": "Missing Limitations section which is required for ACL."}},
    "writing_quality": {{"score": 7, "comment": "Clear prose but some sections lack transitions."}},
    "technical_depth": {{"score": 5, "comment": "Experiments section needs more ablation studies."}},
    "structure":       {{"score": 7, "comment": "Logical flow but Related Work comes too late."}}
  }},
  "recommendations": [
    "Add a Limitations section — this is mandatory for ACL submissions.",
    "Include quantitative baselines in the Experiments section.",
    "Expand the Related Work with recent (2023–2024) citations."
  ],
  "overall_assessment": "The paper presents an interesting approach but is missing key required sections. With revisions it could meet ACL standards."
}}

IMPORTANT:
- section_check must include ALL sections from: {json.dumps(all_sections)}
- quality values: "good" | "fair" | "needs_work" | "missing"
- scores are integers 1–10
- Be specific and actionable in comments and recommendations
- A section is "present" if the headers list contains a matching or similar term
"""

    try:
        resp = client.models.generate_content(model=GEMINI_MODEL, contents=prompt)
        raw = resp.text.strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```\s*$", "", raw)
        result = json.loads(raw)
        result["template_name"]      = reqs["name"]
        result["required_sections"]  = reqs["required"]
        result["recommended_sections"] = reqs["recommended"]
        result["word_count"]         = word_count
        result["sections_found"]     = section_headers
        return result
    except json.JSONDecodeError as e:
        logger.error("Review JSON parse error: %s | raw: %.300s", e, raw)
        raise HTTPException(status_code=500, detail=f"AI returned invalid JSON: {e}")
    except Exception as e:
        logger.error("Review error: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/health")
def health():
    return {"status": "ok", "model": GEMINI_MODEL}